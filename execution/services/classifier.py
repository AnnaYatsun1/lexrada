"""
LLM-based классификация типов контрактов.
25 типов из legal-practice.
"""
from enum import Enum
from typing import Optional



class DocumentType(str, Enum):
    """25 типов контрактов"""
    SUPPLY = "Поставка"
    SERVICE = "Услуги"
    LEASE = "Аренда"
    WORK = "Подряд"
    AGENCY = "Агентский договор"
    LICENSE = "Лицензионный договор"
    MIXED = "Смешанный договор"
    LOAN = "Договор займа"
    PURCHASE = "Купля-продажа"
    DISTRIBUTION = "Дистрибуция/дилерский договор"
    FRANCHISE = "Франчайзинг/коммерческая концессия"
    TRANSPORT = "Перевозка/экспедиция"
    STORAGE = "Хранение"
    COMMISSION = "Комиссия"
    MANDATE = "Поручение"
    DPA = "Договор обработки персональных данных/DPA"
    FRAMEWORK = "Рамочный договор"
    OFFER = "Оферта/пользовательское соглашение"
    CONSTRUCTION = "Договор подряда на строительство"
    EQUIPMENT_LEASE = "Договор аренды оборудования"
    ASSIGNMENT = "Договор уступки права требования"
    EMPLOYMENT = "Трудовой договор"
    NDA = "Соглашение о конфиденциальности/NDA"
    RENTAL = "Договор найма"
    OTHER = "Другое"

DOC_TYPE_TO_RAG_FILTER = {
    DocumentType.EMPLOYMENT: "labor",
    DocumentType.NDA: "nda",
    DocumentType.SERVICE: "services",
    DocumentType.LEASE: "rent",
    DocumentType.RENTAL: "rent",
    DocumentType.EQUIPMENT_LEASE: "rent",
}
def get_rag_filter(doc_type: DocumentType, country: str = "UA") -> dict | None:
    """Возвращает filter для RAG по типу договора. None если RAG-базы для типа нет."""
    rag_doc_type = DOC_TYPE_TO_RAG_FILTER.get(doc_type)
    if rag_doc_type is None:
        return None
    return {"country": country, "doc_type": rag_doc_type}

class DocumentClassifier:
    """LLM-based классификация контрактов"""

    CLASSIFICATION_PROMPT = """Определи ТИП юридического контракта из следующего списка (ТОЛЬКО ТИП, без объяснений):

Поставка, Услуги, Аренда, Подряд, Агентский договор, Лицензионный договор, Смешанный договор, 
Договор займа, Купля-продажа, Дистрибуция/дилерский договор, Франчайзинг, Перевозка/экспедиция, 
Хранение, Комиссия, Поручение, Договор обработки персональных данных, Рамочный договор, 
Оферта/пользовательское соглашение, Договор подряда на строительство, Договор аренды оборудования, 
Договор уступки права требования, Трудовой договор, Соглашение о конфиденциальности, Договор найма, Другое

КОНТРАКТ (первые 2000 символов):
{contract_text}

ОТВЕТ (только тип, например: "Поставка"):
"""

    @classmethod
    def classify(cls, text: str, llm) -> tuple[DocumentType, str]:
        """
        Классифицирует контракт через LLM.
        Возвращает (DocumentType, type_name_ru)
        """
        if not text or len(text) < 100:
            return DocumentType.OTHER, "Другое"

        # Обрезаем текст до 2000 символов
        short_text = text[:2000]
        prompt = cls.CLASSIFICATION_PROMPT.format(contract_text=short_text)

        try:
            response = llm.create(
                model="claude-haiku-4-5",
                max_tokens=50,
                messages=[{"role": "user", "content": prompt}],
                stage="classify",
            )
            
            answer = response.content[0].text.strip()
            
            # Находим совпадение в enum
            for doc_type in DocumentType:
                if doc_type.value.lower() in answer.lower() or answer.lower() in doc_type.value.lower():
                    return doc_type, doc_type.value
            
            # Если не нашли — OTHER
            return DocumentType.OTHER, answer[:50]
            
        except Exception as e:
            print(f"⚠️ Ошибка классификации: {e}")
            return DocumentType.OTHER, "Другое"

    @classmethod
    def get_directive_name(cls, doc_type: DocumentType) -> str:
        """Получить имя файла directive для типа контракта"""
        mapping = {
            DocumentType.SUPPLY: "supply_contract_analysis.md",
            DocumentType.SERVICE: "service_contract_analysis.md",
            DocumentType.LEASE: "lease_contract_analysis.md",
            DocumentType.WORK: "work_contract_analysis.md",
            DocumentType.EMPLOYMENT: "employment_contract_analysis.md",
            DocumentType.NDA: "nda_analysis.md",
            DocumentType.CONSTRUCTION: "construction_contract_analysis.md",
        }
        # Default — есл





from pathlib import Path
from typing import Optional


class DocumentParser:
    """Парсит документы разных форматов в текст"""

    @staticmethod
    def parse_txt(file_path: Path) -> str:
        """Читает текстовый файл"""
        return file_path.read_text(encoding="utf-8")

    @staticmethod
    def parse_docx(file_path: Path) -> str:
        """
        Парсит DOCX файл.
        Требует: pip install python-docx
        """
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx не установлен. Установи: pip install python-docx")

        doc = Document(file_path)
        text_parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Таблицы
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                text_parts.append(" | ".join(row_cells))

        return "\n".join(text_parts)

    @staticmethod
    def parse_pdf(file_path: Path) -> str:
        """
        Парсит PDF файл (текстовый PDF).
        Требует: pip install pdfplumber
        Для отсканированных PDF нужен Tesseract OCR.
        """
        try:
            import pdfplumber
        except ImportError:
            raise ImportError("pdfplumber не установлен. Установи: pip install pdfplumber")

        text_parts = []

        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

                # Пытаемся извлечь таблицы
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            text_parts.append(" | ".join(str(cell) if cell else "" for cell in row))

        return "\n".join(text_parts)

    @staticmethod
    def parse_pdf_ocr(file_path: Path) -> str:
        """
        Парсит отсканированный PDF через OCR.
        Требует: pip install pytesseract pillow
        И установленный Tesseract: https://github.com/UB-Mannheim/tesseract/wiki
        """
        try:
            from pdf2image import convert_from_path
            import pytesseract
            from PIL import Image
        except ImportError:
            raise ImportError(
                "Требуются: pip install pdf2image pytesseract pillow\n"
                "И Tesseract: https://github.com/UB-Mannheim/tesseract/wiki"
            )

        text_parts = []

        # Конвертируем PDF в images
        images = convert_from_path(file_path)

        for image in images:
            text = pytesseract.image_to_string(image, lang="rus+eng")
            if text.strip():
                text_parts.append(text)

        return "\n".join(text_parts)

    @staticmethod
    def parse(file_path: Path, use_ocr: bool = False) -> str:
        """
        Парсит файл любого поддерживаемого формата.
        
        Args:
            file_path: путь к файлу
            use_ocr: использовать OCR для PDF (для отсканированных)
        
        Returns:
            Текст документа
        
        Пример:
            text = DocumentParser.parse(Path("contract.docx"))
            text = DocumentParser.parse(Path("contract.pdf"))
            text = DocumentParser.parse(Path("contract.pdf"), use_ocr=True)  # для отсканированных
        """
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"Файл не найден: {file_path}")

        suffix = file_path.suffix.lower()

        if suffix == ".txt":
            return DocumentParser.parse_txt(file_path)
        elif suffix == ".docx":
            return DocumentParser.parse_docx(file_path)
        elif suffix == ".pdf":
            if use_ocr:
                return DocumentParser.parse_pdf_ocr(file_path)
            else:
                try:
                    return DocumentParser.parse_pdf(file_path)
                except Exception as e:
                    print(f"⚠️ Ошибка при парсинге PDF текстом: {e}")
                    print("💡 Пробую OCR...")
                    return DocumentParser.parse_pdf_ocr(file_path)
        else:
            raise ValueError(f"Неподдерживаемый формат: {suffix}. Поддерживаются: .txt, .pdf, .docx")